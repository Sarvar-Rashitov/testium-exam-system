from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt
from django.utils import timezone
from django.db.models import Avg, Count, Max, Min
from .models import Result
from exams.models import ExamLink
from students.models import Student
import json
import csv
from io import BytesIO
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


@csrf_exempt
def submit_exam_view(request):
    """Submit exam (AJAX endpoint)"""
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            
            student_id = data.get('student_id')
            exam_link_token = data.get('exam_link_token')
            answers = data.get('answers', {})
            time_taken = data.get('time_taken', 0)
            tab_switches = data.get('tab_switches', 0)
            started_at = data.get('started_at')
            
            # Validate
            exam_link = get_object_or_404(ExamLink, unique_token=exam_link_token)
            student = get_object_or_404(Student, pk=student_id)
            exam = exam_link.exam
            
            # Import all question models
            from exams.models import (
                MultipleChoiceSingleQuestion, MultipleChoiceMultipleQuestion,
                TrueFalseNotGivenQuestion, YesNoNotGivenQuestion,
                SentenceCompletionQuestion, ShortAnswerQuestion,
                DiagramLabelingQuestion, SummaryCompletionQuestion,
                NoteCompletionQuestion, TableCompletionQuestion,
                FlowchartCompletionQuestion, MatchingHeadingsQuestion,
                MatchingInformationQuestion, MatchingFeaturesQuestion,
                MatchingSentenceEndingsQuestion
            )
            
            # Map question types to their models
            question_models = {
                'multiple_choice_single': MultipleChoiceSingleQuestion,
                'multiple_choice_multiple': MultipleChoiceMultipleQuestion,
                'true_false_ng': TrueFalseNotGivenQuestion,
                'yes_no_ng': YesNoNotGivenQuestion,
                'sentence_completion': SentenceCompletionQuestion,
                'short_answer': ShortAnswerQuestion,
                'diagram_labeling': DiagramLabelingQuestion,
                'summary_completion': SummaryCompletionQuestion,
                'note_completion': NoteCompletionQuestion,
                'table_completion': TableCompletionQuestion,
                'flowchart_completion': FlowchartCompletionQuestion,
                'matching_headings': MatchingHeadingsQuestion,
                'matching_information': MatchingInformationQuestion,
                'matching_features': MatchingFeaturesQuestion,
                'matching_sentence_endings': MatchingSentenceEndingsQuestion,
            }
            
            # Calculate score
            total_points = 0
            earned_points = 0
            
            # Get all sections with their groups
            sections = exam.sections.all()
            
            for section in sections:
                for group in section.question_groups.all():
                    # Get the appropriate model for this question type
                    question_model = question_models.get(group.question_type)
                    if not question_model:
                        continue
                    
                    # Get all questions from this group
                    questions = question_model.objects.filter(question_group=group).order_by('question_number')
                    
                    for question in questions:
                        # Each question is worth 1 point
                        total_points += 1
                        
                        # Get the answer key format: q_{question_type}_{question_id}
                        answer_key = f"q_{group.question_type}_{question.id}"
                        student_answer = answers.get(answer_key, '')
                        
                        # Check if answer is correct (only for questions with correct answers)
                        if hasattr(question, 'correct_answer') and question.correct_answer:
                            correct_answer = question.correct_answer
                            
                            if isinstance(student_answer, str) and isinstance(correct_answer, str):
                                # For single answer questions
                                if student_answer.strip().upper() == correct_answer.strip().upper():
                                    earned_points += 1
                            elif isinstance(correct_answer, list):
                                # For multiple answer questions
                                student_answers = student_answer.split(',') if isinstance(student_answer, str) else []
                                student_answers = [a.strip().upper() for a in student_answers]
                                correct_answers = [a.strip().upper() for a in correct_answer]
                                if set(student_answers) == set(correct_answers):
                                    earned_points += 1
            
            # If no questions found, set default
            if total_points == 0:
                total_points = 100
            
            percentage = (earned_points / total_points * 100) if total_points > 0 else 0
            
            # Create result
            result = Result.objects.create(
                student=student,
                exam=exam,
                exam_link=exam_link,
                score=earned_points,
                total_points=total_points,
                percentage=percentage,
                answers=answers,
                time_taken=time_taken,
                tab_switches=tab_switches,
                started_at=timezone.datetime.fromisoformat(started_at.replace('Z', '+00:00'))
            )
            
            return JsonResponse({
                'success': True,
                'result_id': result.id,
                'score': earned_points,
                'total_points': total_points,
                'percentage': round(percentage, 2),
                'band_score': result.band_score
            })
            
        except Exception as e:
            return JsonResponse({'success': False, 'error': str(e)}, status=400)
    
    return JsonResponse({'success': False, 'error': 'Invalid request'}, status=400)


def result_view(request, pk):
    """View result (public)"""
    result = get_object_or_404(Result, pk=pk)
    
    # Check if user is admin (organization owner)
    is_admin = request.user.is_authenticated and request.user == result.exam.organization
    
    question_results = []
    template_name = 'results/result_detail.html'
    
    # Only show detailed results to admin
    if is_admin:
        template_name = 'results/result_detail_admin.html'
        
        # Get all sections with their groups
        sections = result.exam.sections.all()
        
        question_counter = 1
        
        # Import all question models
        from exams.models import (
            MultipleChoiceSingleQuestion, MultipleChoiceMultipleQuestion,
            TrueFalseNotGivenQuestion, YesNoNotGivenQuestion,
            SentenceCompletionQuestion, ShortAnswerQuestion,
            DiagramLabelingQuestion, SummaryCompletionQuestion,
            NoteCompletionQuestion, TableCompletionQuestion,
            FlowchartCompletionQuestion, MatchingHeadingsQuestion,
            MatchingInformationQuestion, MatchingFeaturesQuestion,
            MatchingSentenceEndingsQuestion
        )
        
        # Map question types to their models
        question_models = {
            'multiple_choice_single': MultipleChoiceSingleQuestion,
            'multiple_choice_multiple': MultipleChoiceMultipleQuestion,
            'true_false_ng': TrueFalseNotGivenQuestion,
            'yes_no_ng': YesNoNotGivenQuestion,
            'sentence_completion': SentenceCompletionQuestion,
            'short_answer': ShortAnswerQuestion,
            'diagram_labeling': DiagramLabelingQuestion,
            'summary_completion': SummaryCompletionQuestion,
            'note_completion': NoteCompletionQuestion,
            'table_completion': TableCompletionQuestion,
            'flowchart_completion': FlowchartCompletionQuestion,
            'matching_headings': MatchingHeadingsQuestion,
            'matching_information': MatchingInformationQuestion,
            'matching_features': MatchingFeaturesQuestion,
            'matching_sentence_endings': MatchingSentenceEndingsQuestion,
        }
        
        for section in sections:
            for group in section.question_groups.all():
                # Get the appropriate model for this question type
                question_model = question_models.get(group.question_type)
                if not question_model:
                    continue
                
                # Get all questions from this group
                questions = question_model.objects.filter(question_group=group).order_by('question_number')
                
                for question in questions:
                    # Get the answer key format: q_{question_type}_{question_id}
                    answer_key = f"q_{group.question_type}_{question.id}"
                    student_answer = result.answers.get(answer_key, '')
                    
                    # Check if answer is correct (only for questions with correct answers)
                    is_correct = False
                    correct_answer = None
                    
                    if hasattr(question, 'correct_answer') and question.correct_answer:
                        correct_answer = question.correct_answer
                        if isinstance(student_answer, str) and isinstance(correct_answer, str):
                            # For single answer questions
                            is_correct = student_answer.strip().upper() == correct_answer.strip().upper()
                        elif isinstance(correct_answer, list):
                            # For multiple answer questions
                            student_answers = student_answer.split(',') if isinstance(student_answer, str) else []
                            student_answers = [a.strip().upper() for a in student_answers]
                            correct_answers = [a.strip().upper() for a in correct_answer]
                            is_correct = set(student_answers) == set(correct_answers)
                    
                    # Get question text - different question types have different field names
                    question_text = ''
                    if hasattr(question, 'statement'):
                        question_text = question.statement
                    elif hasattr(question, 'question_text'):
                        question_text = question.question_text
                    elif hasattr(question, 'text'):
                        question_text = question.text
                    elif hasattr(question, 'sentence'):
                        question_text = question.sentence
                    
                    question_results.append({
                        'question': question,
                        'question_text': question_text,
                        'student_answer': student_answer,
                        'is_correct': is_correct,
                        'correct_answer': correct_answer,
                        'counter': question_counter,
                        'question_type': group.question_type
                    })
                    question_counter += 1
    
    # Count incorrect answers
    incorrect_count = sum(1 for item in question_results if not item['is_correct'] and item['correct_answer'])
    
    context = {
        'result': result,
        'question_results': question_results,
        'is_admin': is_admin,
        'time_taken_minutes': result.time_taken // 60 if result.time_taken else 0,
        'incorrect_count': incorrect_count
    }
    
    return render(request, template_name, context)


@login_required
def results_list_view(request):
    """List all results (admin)"""
    from django.core.paginator import Paginator
    from django.db.models import Q
    
    results = Result.objects.filter(exam__organization=request.user).select_related('student', 'exam', 'student__teacher')
    
    # Search functionality - search in first_name, last_name, phone, exam title, teacher name
    search_query = request.GET.get('search', '').strip()
    if search_query:
        results = results.filter(
            Q(student__first_name__icontains=search_query) |
            Q(student__last_name__icontains=search_query) |
            Q(student__phone__icontains=search_query) |
            Q(exam__title__icontains=search_query) |
            Q(student__teacher__first_name__icontains=search_query) |
            Q(student__teacher__last_name__icontains=search_query)
        )
    
    # Filter by exam
    exam_id = request.GET.get('exam_id')
    if exam_id:
        results = results.filter(exam_id=exam_id)
    
    # Filter by student type
    student_type = request.GET.get('student_type')
    if student_type:
        results = results.filter(student__student_type=student_type)
    
    # Filter by teacher
    teacher_id = request.GET.get('teacher_id')
    if teacher_id:
        results = results.filter(student__teacher_id=teacher_id)
    
    # Order by latest first
    results = results.order_by('-completed_at')
    
    # Pagination
    paginator = Paginator(results, 10)  # 10 results per page
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)
    
    # Get exams and teachers for filters
    from exams.models import Exam
    from organizations.models import Teacher
    exams = Exam.objects.filter(organization=request.user)
    teachers = Teacher.objects.filter(organization=request.user, is_active=True)
    
    context = {
        'results': page_obj,
        'page_obj': page_obj,
        'exams': exams,
        'teachers': teachers,
        'selected_exam_id': exam_id,
        'selected_student_type': student_type,
        'selected_teacher_id': teacher_id,
        'search_query': search_query,
        'total_results': paginator.count
    }
    
    return render(request, 'results/results_list.html', context)


@login_required
def analytics_view(request):
    """Analytics dashboard"""
    exam_id = request.GET.get('exam_id')
    results = Result.objects.filter(exam__organization=request.user)
    
    if exam_id:
        results = results.filter(exam_id=exam_id)
    
    # Calculate analytics
    analytics = results.aggregate(
        total_students=Count('student', distinct=True),
        total_attempts=Count('id'),
        avg_score=Avg('score'),
        avg_percentage=Avg('percentage'),
        max_score=Max('score'),
        min_score=Min('score'),
        avg_time_taken=Avg('time_taken')
    )
    
    # Score distribution
    score_distribution = [
        {'range': '0-20%', 'count': results.filter(percentage__lt=20).count()},
        {'range': '20-40%', 'count': results.filter(percentage__gte=20, percentage__lt=40).count()},
        {'range': '40-60%', 'count': results.filter(percentage__gte=40, percentage__lt=60).count()},
        {'range': '60-80%', 'count': results.filter(percentage__gte=60, percentage__lt=80).count()},
        {'range': '80-100%', 'count': results.filter(percentage__gte=80).count()},
    ]
    
    # Band score distribution
    band_distribution = {}
    for result in results:
        band = result.band_score
        band_distribution[band] = band_distribution.get(band, 0) + 1
    
    # Sort by band score
    band_distribution = dict(sorted(band_distribution.items()))
    
    # Results over time (last 30 days)
    from datetime import timedelta
    thirty_days_ago = timezone.now() - timedelta(days=30)
    recent_results = results.filter(completed_at__gte=thirty_days_ago)
    
    # Group by date
    results_by_date = {}
    for result in recent_results:
        date_key = result.completed_at.strftime('%Y-%m-%d')
        if date_key not in results_by_date:
            results_by_date[date_key] = {'count': 0, 'total_percentage': 0}
        results_by_date[date_key]['count'] += 1
        results_by_date[date_key]['total_percentage'] += float(result.percentage)
    
    # Calculate average for each day
    timeline_data = []
    for date_key in sorted(results_by_date.keys()):
        data = results_by_date[date_key]
        avg_percentage = data['total_percentage'] / data['count']
        timeline_data.append({
            'date': date_key,
            'count': data['count'],
            'avg_percentage': round(avg_percentage, 2)
        })
    
    # Student type distribution
    from students.models import Student
    student_type_dist = results.values('student__student_type').annotate(
        count=Count('id')
    ).order_by('student__student_type')
    
    # Top students - limit to 5
    top_students = results.order_by('-percentage', '-completed_at')[:5]
    
    # Get exams for filter
    from exams.models import Exam
    exams = Exam.objects.filter(organization=request.user)
    
    context = {
        'analytics': analytics,
        'score_distribution': score_distribution,
        'band_distribution': band_distribution,
        'timeline_data': timeline_data,
        'student_type_dist': student_type_dist,
        'top_students': top_students,
        'exams': exams,
        'selected_exam_id': exam_id
    }
    
    return render(request, 'results/analytics.html', context)


@login_required
def export_csv_view(request):
    """Export results to CSV"""
    exam_id = request.GET.get('exam_id')
    results = Result.objects.filter(exam__organization=request.user).select_related('student', 'exam')
    
    if exam_id:
        results = results.filter(exam_id=exam_id)
    
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = 'attachment; filename="natijalar.csv"'
    
    # Add BOM for Excel UTF-8 support
    response.write('\ufeff')
    
    writer = csv.writer(response)
    writer.writerow(['Talaba', 'Telefon', 'Telegram', 'Imtihon', 'Ball', 'Foiz', 'IELTS Band', 'Vaqt (daq)', 'Sana'])
    
    for result in results:
        writer.writerow([
            result.student.full_name,
            result.student.phone,
            result.student.telegram_username,
            result.exam.title,
            f"{result.score}/{result.total_points}",
            f"{result.percentage:.1f}%",
            result.band_score,
            result.time_taken // 60 if result.time_taken else 0,
            result.completed_at.strftime('%Y-%m-%d %H:%M')
        ])
    
    return response


@login_required
def export_excel_view(request):
    """Export results to Excel (XLSX)"""
    exam_id = request.GET.get('exam_id')
    results = Result.objects.filter(exam__organization=request.user).select_related('student', 'exam', 'student__teacher')
    
    if exam_id:
        results = results.filter(exam_id=exam_id)
    
    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Natijalar"
    
    # Define styles
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    border_style = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    
    cell_alignment = Alignment(horizontal="left", vertical="center")
    center_alignment = Alignment(horizontal="center", vertical="center")
    
    # Headers
    headers = ['#', 'Talaba', 'Telefon', 'Telegram', 'Turi', "O'qituvchi", 'Imtihon', 'Ball', 'Foiz', 'IELTS Band', 'Vaqt (daq)', 'Tab Switch', 'Sana']
    ws.append(headers)
    
    # Style header row
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border_style
    
    # Add data
    for idx, result in enumerate(results, 1):
        student_type = 'Muassasa' if result.student.student_type == 'institution' else 'Tashqi'
        teacher_name = result.student.teacher.full_name if result.student.teacher else '-'
        
        row_data = [
            idx,
            result.student.full_name,
            result.student.phone,
            result.student.telegram_username or '-',
            student_type,
            teacher_name,
            result.exam.title,
            f"{result.score}/{result.total_points}",
            f"{result.percentage:.1f}%",
            result.band_score,
            result.time_taken // 60 if result.time_taken else 0,
            result.tab_switches,
            result.completed_at.strftime('%d.%m.%Y %H:%M')
        ]
        ws.append(row_data)
        
        # Style data cells
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=idx + 1, column=col_num)
            cell.border = border_style
            
            # Center alignment for specific columns
            if col_num in [1, 5, 8, 9, 10, 11, 12]:  # #, Turi, Ball, Foiz, Band, Vaqt, Tab Switch
                cell.alignment = center_alignment
            else:
                cell.alignment = cell_alignment
    
    # Adjust column widths
    column_widths = [5, 25, 15, 15, 12, 20, 30, 12, 10, 12, 12, 12, 18]
    for idx, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    
    # Set row height for header
    ws.row_dimensions[1].height = 25
    
    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    # Create response
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="natijalar.xlsx"'
    
    return response


@login_required
def export_word_view(request):
    """Export results to Word (DOCX)"""
    exam_id = request.GET.get('exam_id')
    results = Result.objects.filter(exam__organization=request.user).select_related('student', 'exam', 'student__teacher')
    
    if exam_id:
        results = results.filter(exam_id=exam_id)
    
    # Create document
    doc = Document()
    
    # Add title
    title = doc.add_heading('Imtihon Natijalari', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add organization info
    org_name = request.user.full_name if hasattr(request.user, 'full_name') else request.user.username
    org_para = doc.add_paragraph(org_name)
    org_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    org_para.runs[0].font.size = Pt(14)
    org_para.runs[0].font.bold = True
    
    # Add date
    date_para = doc.add_paragraph(f"Sana: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_para.runs[0].font.size = Pt(11)
    date_para.runs[0].font.color.rgb = RGBColor(107, 114, 128)
    
    doc.add_paragraph()  # Empty line
    
    # Create table
    table = doc.add_table(rows=1, cols=13)
    table.style = 'Light Grid Accent 1'
    
    # Header row
    header_cells = table.rows[0].cells
    headers = ['#', 'Talaba', 'Telefon', 'Telegram', 'Turi', "O'qituvchi", 'Imtihon', 'Ball', 'Foiz', 'Band', 'Vaqt', 'Tab', 'Sana']
    
    for idx, header in enumerate(headers):
        cell = header_cells[idx]
        cell.text = header
        # Style header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(79, 70, 229)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add data rows
    for idx, result in enumerate(results, 1):
        row_cells = table.add_row().cells
        student_type = 'Muassasa' if result.student.student_type == 'institution' else 'Tashqi'
        teacher_name = result.student.teacher.full_name if result.student.teacher else '-'
        
        data = [
            str(idx),
            result.student.full_name,
            result.student.phone,
            result.student.telegram_username or '-',
            student_type,
            teacher_name,
            result.exam.title,
            f"{result.score}/{result.total_points}",
            f"{result.percentage:.1f}%",
            str(result.band_score),
            str(result.time_taken // 60 if result.time_taken else 0),
            str(result.tab_switches),
            result.completed_at.strftime('%d.%m.%Y')
        ]
        
        for cell_idx, value in enumerate(data):
            cell = row_cells[cell_idx]
            cell.text = str(value)  # Ensure all values are strings
            # Style cell
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    
    # Save to BytesIO
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    
    # Create response
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    response['Content-Disposition'] = 'attachment; filename="natijalar.docx"'
    
    return response


@login_required
def export_pdf_view(request):
    """Export results to PDF"""
    exam_id = request.GET.get('exam_id')
    results = Result.objects.filter(exam__organization=request.user).select_related('student', 'exam', 'student__teacher')
    
    if exam_id:
        results = results.filter(exam_id=exam_id)
    
    # Create PDF in memory
    buffer = BytesIO()
    p = canvas.Canvas(buffer, pagesize=landscape(A4))
    width, height = landscape(A4)
    
    # Organization name
    org_name = request.user.full_name if hasattr(request.user, 'full_name') else request.user.username
    
    # Title
    p.setFont("Helvetica-Bold", 18)
    p.setFillColor(colors.HexColor('#1f2937'))
    p.drawCentredString(width/2, height-2*cm, "Imtihon Natijalari")
    
    # Organization
    p.setFont("Helvetica-Bold", 14)
    p.setFillColor(colors.HexColor('#4F46E5'))
    p.drawCentredString(width/2, height-2.8*cm, org_name)
    
    # Date
    p.setFont("Helvetica", 10)
    p.setFillColor(colors.HexColor('#6b7280'))
    p.drawCentredString(width/2, height-3.5*cm, f"Sana: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    
    # Table headers
    y_position = height - 5*cm
    x_start = 1.5*cm
    
    # Column widths
    col_widths = [1*cm, 4*cm, 3*cm, 3*cm, 2.5*cm, 3.5*cm, 5*cm, 2*cm, 2*cm, 2*cm, 1.5*cm, 1.5*cm, 3*cm]
    
    # Draw header background
    p.setFillColor(colors.HexColor('#4F46E5'))
    p.rect(x_start, y_position-0.3*cm, sum(col_widths), 0.8*cm, fill=True, stroke=False)
    
    # Header text
    headers = ['#', 'Talaba', 'Telefon', 'Telegram', 'Turi', "O'qituvchi", 'Imtihon', 'Ball', 'Foiz', 'Band', 'Vaqt', 'Tab', 'Sana']
    p.setFont("Helvetica-Bold", 8)
    p.setFillColor(colors.white)
    
    x_pos = x_start
    for idx, header in enumerate(headers):
        p.drawString(x_pos + 0.1*cm, y_position, header)
        x_pos += col_widths[idx]
    
    # Data rows
    y_position -= 1*cm
    p.setFont("Helvetica", 7)
    
    row_count = 0
    max_rows_per_page = 15
    
    for idx, result in enumerate(results, 1):
        if row_count >= max_rows_per_page:
            # New page
            p.showPage()
            y_position = height - 2*cm
            row_count = 0
            
            # Redraw header on new page
            p.setFillColor(colors.HexColor('#4F46E5'))
            p.rect(x_start, y_position-0.3*cm, sum(col_widths), 0.8*cm, fill=True, stroke=False)
            p.setFont("Helvetica-Bold", 8)
            p.setFillColor(colors.white)
            x_pos = x_start
            for h_idx, header in enumerate(headers):
                p.drawString(x_pos + 0.1*cm, y_position, header)
                x_pos += col_widths[h_idx]
            y_position -= 1*cm
            p.setFont("Helvetica", 7)
        
        # Alternate row colors
        if idx % 2 == 0:
            p.setFillColor(colors.HexColor('#f9fafb'))
            p.rect(x_start, y_position-0.3*cm, sum(col_widths), 0.7*cm, fill=True, stroke=False)
        
        p.setFillColor(colors.HexColor('#1f2937'))
        
        student_type = 'Muassasa' if result.student.student_type == 'institution' else 'Tashqi'
        teacher_name = result.student.teacher.full_name if result.student.teacher else '-'
        
        data = [
            str(idx),
            result.student.full_name[:20],  # Truncate long names
            result.student.phone,
            (result.student.telegram_username or '-')[:15],
            student_type,
            teacher_name[:18],
            result.exam.title[:25],
            f"{result.score}/{result.total_points}",
            f"{result.percentage:.1f}%",
            result.band_score,
            str(result.time_taken // 60 if result.time_taken else 0),
            str(result.tab_switches),
            result.completed_at.strftime('%d.%m.%Y')
        ]
        
        x_pos = x_start
        for d_idx, value in enumerate(data):
            p.drawString(x_pos + 0.1*cm, y_position, str(value))
            x_pos += col_widths[d_idx]
        
        y_position -= 0.7*cm
        row_count += 1
    
    # Footer
    p.setFont("Helvetica-Oblique", 8)
    p.setFillColor(colors.HexColor('#9ca3af'))
    p.drawCentredString(width/2, 1*cm, f"Jami {len(results)} ta natija")
    
    p.showPage()
    p.save()
    
    buffer.seek(0)
    
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="natijalar.pdf"'
    
    return response



def download_certificate_view(request, pk):
    """Generate and download IELTS-style certificate"""
    result = get_object_or_404(Result, pk=pk)
    
    # Create PDF in memory
    buffer = BytesIO()
    
    # Create PDF with landscape A4 size
    p = canvas.Canvas(buffer, pagesize=landscape(A4))
    width, height = landscape(A4)
    
    # Background color
    p.setFillColor(colors.HexColor('#f8f9fc'))
    p.rect(0, 0, width, height, fill=True, stroke=False)
    
    # Border
    p.setStrokeColor(colors.HexColor('#4F46E5'))
    p.setLineWidth(3)
    p.rect(1*cm, 1*cm, width-2*cm, height-2*cm, fill=False, stroke=True)
    
    # Inner border
    p.setStrokeColor(colors.HexColor('#6366F1'))
    p.setLineWidth(1)
    p.rect(1.3*cm, 1.3*cm, width-2.6*cm, height-2.6*cm, fill=False, stroke=True)
    
    # Header - Organization name
    p.setFillColor(colors.HexColor('#1f2937'))
    p.setFont("Helvetica-Bold", 24)
    org_name = result.exam.organization.full_name if hasattr(result.exam.organization, 'full_name') else result.exam.organization.username
    p.drawCentredString(width/2, height-3*cm, org_name.upper())
    
    # Certificate title
    p.setFillColor(colors.HexColor('#4F46E5'))
    p.setFont("Helvetica-Bold", 36)
    p.drawCentredString(width/2, height-5*cm, "CERTIFICATE OF ACHIEVEMENT")
    
    # Subtitle
    p.setFillColor(colors.HexColor('#6b7280'))
    p.setFont("Helvetica", 14)
    p.drawCentredString(width/2, height-6*cm, "This is to certify that")
    
    # Student name
    p.setFillColor(colors.HexColor('#1f2937'))
    p.setFont("Helvetica-Bold", 32)
    p.drawCentredString(width/2, height-8*cm, result.student.full_name.upper())
    
    # Line under name
    p.setStrokeColor(colors.HexColor('#4F46E5'))
    p.setLineWidth(2)
    p.line(width/2-8*cm, height-8.5*cm, width/2+8*cm, height-8.5*cm)
    
    # Achievement text
    p.setFillColor(colors.HexColor('#6b7280'))
    p.setFont("Helvetica", 14)
    p.drawCentredString(width/2, height-9.5*cm, "has successfully completed")
    
    # Exam name
    p.setFillColor(colors.HexColor('#1f2937'))
    p.setFont("Helvetica-Bold", 20)
    p.drawCentredString(width/2, height-11*cm, result.exam.title)
    
    # Results box
    box_y = height - 14*cm
    box_height = 2.5*cm
    
    # Score box
    p.setFillColor(colors.HexColor('#eff6ff'))
    p.rect(width/2-10*cm, box_y, 6*cm, box_height, fill=True, stroke=False)
    p.setStrokeColor(colors.HexColor('#3b82f6'))
    p.setLineWidth(1)
    p.rect(width/2-10*cm, box_y, 6*cm, box_height, fill=False, stroke=True)
    
    p.setFillColor(colors.HexColor('#1f2937'))
    p.setFont("Helvetica", 12)
    p.drawCentredString(width/2-7*cm, box_y+1.8*cm, "Score")
    p.setFont("Helvetica-Bold", 20)
    p.drawCentredString(width/2-7*cm, box_y+0.6*cm, f"{result.score}/{result.total_points}")
    
    # Percentage box
    p.setFillColor(colors.HexColor('#f0fdf4'))
    p.rect(width/2-3.5*cm, box_y, 6*cm, box_height, fill=True, stroke=False)
    p.setStrokeColor(colors.HexColor('#22c55e'))
    p.setLineWidth(1)
    p.rect(width/2-3.5*cm, box_y, 6*cm, box_height, fill=False, stroke=True)
    
    p.setFillColor(colors.HexColor('#1f2937'))
    p.setFont("Helvetica", 12)
    p.drawCentredString(width/2-0.5*cm, box_y+1.8*cm, "Percentage")
    p.setFont("Helvetica-Bold", 20)
    p.drawCentredString(width/2-0.5*cm, box_y+0.6*cm, f"{result.percentage:.1f}%")
    
    # Band score box
    p.setFillColor(colors.HexColor('#fef3c7'))
    p.rect(width/2+3*cm, box_y, 6*cm, box_height, fill=True, stroke=False)
    p.setStrokeColor(colors.HexColor('#f59e0b'))
    p.setLineWidth(1)
    p.rect(width/2+3*cm, box_y, 6*cm, box_height, fill=False, stroke=True)
    
    p.setFillColor(colors.HexColor('#1f2937'))
    p.setFont("Helvetica", 12)
    p.drawCentredString(width/2+6*cm, box_y+1.8*cm, "IELTS Band")
    p.setFont("Helvetica-Bold", 20)
    p.drawCentredString(width/2+6*cm, box_y+0.6*cm, f"{result.band_score}")
    
    # Date
    p.setFillColor(colors.HexColor('#6b7280'))
    p.setFont("Helvetica", 11)
    date_str = result.completed_at.strftime("%d %B %Y")
    p.drawString(3*cm, 2.5*cm, f"Date: {date_str}")
    
    # Certificate ID
    p.drawRightString(width-3*cm, 2.5*cm, f"Certificate ID: {result.id:06d}")
    
    # Footer text
    p.setFont("Helvetica-Oblique", 10)
    p.setFillColor(colors.HexColor('#9ca3af'))
    p.drawCentredString(width/2, 1.8*cm, "This certificate verifies the completion and achievement in the examination")
    
    # Finalize PDF
    p.showPage()
    p.save()
    
    # Get PDF from buffer
    buffer.seek(0)
    
    # Create response
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    filename = f"certificate_{result.student.full_name.replace(' ', '_')}_{result.exam.title.replace(' ', '_')}.pdf"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    return response
